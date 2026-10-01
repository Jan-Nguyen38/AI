"""Turn files and web pages into plain text, keeping page numbers for PDFs."""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass
from pathlib import Path

import httpx
from bs4 import BeautifulSoup

log = logging.getLogger("rag.loaders")

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".html", ".htm"}


@dataclass
class Section:
    """A piece of a document's text. `page` is set for PDFs (1-based)."""

    text: str
    page: int | None = None


def load_bytes(name: str, data: bytes) -> list[Section]:
    ext = Path(name).suffix.lower()
    log.debug("Loading %s (%d bytes)", name, len(data))
    if ext == ".pdf":
        return _load_pdf(data)
    if ext == ".docx":
        return _load_docx(data)
    if ext in {".html", ".htm"}:
        return [Section(_html_to_text(data.decode("utf-8", errors="replace")))]
    if ext in {".txt", ".md"}:
        return [Section(data.decode("utf-8", errors="replace"))]
    raise ValueError(f"Unsupported file type: {ext or name}. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}")


def load_path(path: Path) -> list[Section]:
    return load_bytes(path.name, path.read_bytes())


def load_url(url: str) -> tuple[str, list[Section]]:
    """Fetch a web page and return (title, sections)."""
    log.info("Fetching %s", url)
    resp = httpx.get(url, follow_redirects=True, timeout=30, headers={"User-Agent": "rag-chatbot/0.1"})
    resp.raise_for_status()
    if "application/pdf" in resp.headers.get("content-type", ""):
        return url, _load_pdf(resp.content)
    soup = BeautifulSoup(resp.text, "html.parser")
    title = (soup.title.string or url).strip() if soup.title else url
    return title, [Section(_html_to_text(resp.text))]


def _load_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(text, page=i))
    log.debug("PDF has %d pages, %d with text", len(reader.pages), len(sections))
    if not sections:
        raise ValueError("No text found in PDF. Scanned PDFs need OCR, which this version doesn't do yet.")
    return sections


def _load_docx(data: bytes) -> list[Section]:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Section("\n\n".join(parts))]


def _html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "svg"]):
        tag.decompose()
    lines = (line.strip() for line in soup.get_text("\n").splitlines())
    return "\n".join(line for line in lines if line)
